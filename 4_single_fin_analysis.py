import os
import re
import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from company_config import PROJECT_DIR, get_company_name

# 회사는 .env의 COMPANY_NAME에서 설정합니다. 분석 기간은 years를 수정하세요.

# ==========================================
# 1. 파일 경로 설정 및 데이터 로드
# ==========================================
COMPANY_NAME = get_company_name()
file_dir = str(PROJECT_DIR)
input_file = os.path.join(file_dir, f"{COMPANY_NAME}_재무제표_3단계_완전정제_리포트.xlsx")
output_file = os.path.join(file_dir, f"{COMPANY_NAME}_재무분석_결과리포트_최종.xlsx")

# Excel 데이터 로드
df = pd.read_excel(input_file, sheet_name="3_최종_정제통합본")

# 계정과목명 및 account_id 전처리 (공백 제거 및 문자열 변환)
if "계정과목명" in df.columns:
  df["계정과목명_정규화"] = (
      df["계정과목명"].astype(str).str.replace(r"\s+", "", regex=True)
  )
else:
  df["계정과목명_정규화"] = ""

if "account_id" in df.columns:
  df["account_id_str"] = df["account_id"].astype(str).str.strip()
else:
  df["account_id_str"] = ""

years = [f"{y}년" for y in range(2016, 2026)]


# ==========================================
# 2. 초범용 강력 계정 추출 함수 (우선순위 기반 다중 검색)
# ==========================================
def extract_financial_series_ultra(
    df, target_ids=[], exact_names=[], contain_names=[], exclude_names=[]
):
  """1순위: account_id 매칭

  2순위: 계정과목명 '완전 일치' 매칭 3순위: 계정과목명 '포함' 매칭 (제외 단어 필터링)
  """
  result = {}

  exact_clean = [nm.replace(" ", "") for nm in exact_names]
  contain_clean = [nm.replace(" ", "") for nm in contain_names]
  exclude_clean = [ex.replace(" ", "") for ex in exclude_names]

  for year_str in years:
    df_year = df[df["연도"] == year_str].copy()
    found_val = np.nan

    # [1단계] 표준 account_id 매칭
    if target_ids and "account_id_str" in df_year.columns:
      matched = df_year[
          df_year["account_id_str"].isin(target_ids)
          & df_year["당기금액"].notna()
      ]
      if not matched.empty:
        found_val = matched["당기금액"].values[0]

    # [2단계] 계정과목명 '완전 일치' 탐색
    if pd.isna(found_val) and exact_clean:
      matched = df_year[
          df_year["계정과목명_정규화"].isin(exact_clean)
          & df_year["당기금액"].notna()
      ]
      if not matched.empty:
        found_val = matched["당기금액"].values[0]

    # [3단계] 계정과목명 '포함' 탐색 (제외 단어 검증)
    if pd.isna(found_val) and contain_clean:
      for kw in contain_clean:
        matched = df_year[
            df_year["계정과목명_정규화"].str.contains(kw, regex=False, na=False)
            & df_year["당기금액"].notna()
        ]

        for ex in exclude_clean:
          matched = matched[
              ~matched["계정과목명_정규화"].str.contains(ex, regex=False)
          ]

        if not matched.empty:
          found_val = matched["당기금액"].values[0]
          break

    result[year_str] = found_val

  return pd.Series(result)


# ==========================================
# 3. 계정별 매핑 정의 및 수치 추출
# ==========================================

# 1) 매출액 (DART 표준 ID 및 표기법 완벽 포획)
sales = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_Revenue",
        "ifrs_Revenue",
        "dart_Revenue",
        "ifrs-full_SalesRevenueNet",
        "dart_OperatingRevenue",
    ],
    exact_names=["매출액", "수익(매출액)", "영업수익", "매출", "수익"],
    contain_names=["매출액", "영업수익", "수익(매출액)"],
    exclude_names=[
        "매출채권",
        "매출원가",
        "매출총이익",
        "기타수익",
        "금융수익",
        "지분법수익",
    ],
)

# 2) 영업이익
op_inc = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_OperatingProfitLoss",
        "ifrs_OperatingProfitLoss",
        "dart_OperatingIncomeLoss",
        "dart_OperatingProfit",
    ],
    exact_names=["영업이익", "영업손익", "영업이익(손실)", "영업손익(손실)"],
    contain_names=["영업이익", "영업손익"],
)

# 3) 당기순이익
net_inc = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_ProfitLoss",
        "ifrs_ProfitLoss",
        "dart_ProfitLoss",
        "ifrs-full_ProfitLossAttributableToOwnersOfParent",
    ],
    exact_names=[
        "당기순이익",
        "당기순손실",
        "순이익",
        "당기순이익(손실)",
        "반기순이익",
        "분기순이익",
    ],
    contain_names=["당기순이익", "당기순손실"],
    exclude_names=["총포괄", "기타포괄", "지배기업"],
)

# 4) 자산총계
assets = extract_financial_series_ultra(
    df,
    target_ids=["ifrs-full_Assets", "ifrs_Assets", "dart_Assets"],
    exact_names=["자산총계", "자산합계"],
    contain_names=["자산총계"],
    exclude_names=["유동자산", "비유동자산", "순자산"],
)

# 5) 자본총계
equity = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_Equity",
        "ifrs_Equity",
        "dart_Equity",
        "ifrs-full_EquityAttributableToOwnersOfParent",
    ],
    exact_names=["자본총계", "자본합계"],
    contain_names=["자본총계"],
    exclude_names=["유동자본", "납입자본"],
)

# 6) 부채총계
liab = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_Liabilities",
        "ifrs_Liabilities",
        "dart_Liabilities",
    ],
    exact_names=["부채총계", "부채합계"],
    contain_names=["부채총계"],
    exclude_names=["유동부채", "비유동부채"],
)

# 7) 유동자산
cur_assets = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_CurrentAssets",
        "ifrs_CurrentAssets",
        "dart_CurrentAssets",
    ],
    exact_names=["유동자산"],
    contain_names=["유동자산"],
    exclude_names=["비유동자산"],
)

# 8) 유동부채
cur_liab = extract_financial_series_ultra(
    df,
    target_ids=[
        "ifrs-full_CurrentLiabilities",
        "ifrs_CurrentLiabilities",
        "dart_CurrentLiabilities",
    ],
    exact_names=["유동부채"],
    contain_names=["유동부채"],
    exclude_names=["비유동부채"],
)

# ==========================================
# 4. Fallback (보완 연산) 및 주요 재무비율 계산
# ==========================================
# 자산 = 부채 + 자본 방정식 이용 결측치 보완
assets = assets.fillna(liab + equity)
equity = equity.fillna(assets - liab)
liab = liab.fillna(assets - equity)

# 수익성/안정성/성장성 재무비율 산출
op_margin = op_inc / sales
net_margin = net_inc / sales
roa = net_inc / assets
roe = net_inc / equity

debt_ratio = liab / equity
current_ratio = cur_assets / cur_liab

sales_growth = sales.pct_change()
op_growth = op_inc.pct_change()
asset_growth = assets.pct_change()

# ==========================================
# 5. Openpyxl 기반 서식 지정 엑셀 리포트 생성
# ==========================================
ratio_rows = [
    ("규모 및 성과", "매출액 (억원)", (sales / 1e8).values, "#,##0"),
    ("규모 및 성과", "영업이익 (억원)", (op_inc / 1e8).values, "#,##0"),
    ("규모 및 성과", "당기순이익 (억원)", (net_inc / 1e8).values, "#,##0"),
    ("규모 및 성과", "자산총계 (억원)", (assets / 1e8).values, "#,##0"),
    ("규모 및 성과", "자본총계 (억원)", (equity / 1e8).values, "#,##0"),
    ("수익성 지표", "영업이익률 (%)", op_margin.values, "0.0%"),
    ("수익성 지표", "순이익률 (%)", net_margin.values, "0.0%"),
    ("수익성 지표", "ROA (%)", roa.values, "0.0%"),
    ("수익성 지표", "ROE (%)", roe.values, "0.0%"),
    ("안정성 지표", "부채비율 (%)", debt_ratio.values, "0.0%"),
    ("안정성 지표", "유동비율 (%)", current_ratio.values, "0.0%"),
    ("성장성 지표", "매출액증가율 (%)", sales_growth.values, "0.0%"),
    ("성장성 지표", "영업이익증가율 (%)", op_growth.values, "0.0%"),
    ("성장성 지표", "총자산증가율 (%)", asset_growth.values, "0.0%"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "재무제표_종합분석"
ws.views.sheetView[0].showGridLines = True

# 서식 스타일 지정
font_family = "맑은 고딕"
title_font = Font(name=font_family, size=15, bold=True, color="1B365D")
subtitle_font = Font(name=font_family, size=10, italic=True, color="555555")
section_font = Font(name=font_family, size=10, bold=True, color="1B365D")
header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
regular_font = Font(name=font_family, size=10)

header_fill = PatternFill(
    start_color="1B365D", end_color="1B365D", fill_type="solid"
)
cat_fill = PatternFill(
    start_color="F0F4F8", end_color="F0F4F8", fill_type="solid"
)

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# 타이틀 작성
ws.cell(
    row=2,
    column=2,
    value=f"{COMPANY_NAME} 재무제표 종합 분석 리포트 ({years[0].removesuffix('년')}~{years[-1].removesuffix('년')})",
).font = title_font
ws.cell(
    row=3, column=2, value="자동 계정 탐색 및 비율 연산 완료"
).font = subtitle_font

# 헤더 작성
headers = ["분석 구분", "세부 재무지표"] + years
ws.row_dimensions[5].height = 24
for c_idx, h in enumerate(headers, start=2):
  cell = ws.cell(row=5, column=c_idx, value=h)
  cell.font = header_font
  cell.fill = header_fill
  cell.alignment = align_center

# 데이터 채우기
start_row = 6
for item in ratio_rows:
  cat, metric_name, values, num_format = item
  ws.row_dimensions[start_row].height = 20

  cell_cat = ws.cell(row=start_row, column=2, value=cat)
  cell_cat.font = section_font
  cell_cat.border = thin_border
  cell_cat.alignment = align_center
  cell_cat.fill = cat_fill

  cell_metric = ws.cell(row=start_row, column=3, value=metric_name)
  cell_metric.font = regular_font
  cell_metric.border = thin_border
  cell_metric.alignment = align_left

  for idx, val in enumerate(values):
    col_idx = 4 + idx
    cell_val = ws.cell(row=start_row, column=col_idx)
    cell_val.font = regular_font
    cell_val.border = thin_border
    cell_val.alignment = align_right

    if pd.isna(val) or np.isnan(val):
      cell_val.value = "-"
      cell_val.alignment = align_center
    else:
      cell_val.value = float(val)
      cell_val.number_format = num_format

  start_row += 1

# 열 너비 설정
for col in ws.columns:
  col_letter = get_column_letter(col[0].column)
  if col_letter == "A":
    ws.column_dimensions[col_letter].width = 3
    continue
  ws.column_dimensions[col_letter].width = 15

ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 22

# 파일 저장
wb.save(output_file)
print(f"✅ 파일 생성이 완료되었습니다: {output_file}")

